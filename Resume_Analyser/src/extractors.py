import io
from pathlib import Path
from typing import Union, BinaryIO

def extract_text_from_pdf(file_source: Union[str, Path, BinaryIO]) -> str:
    """Extract clean text content from a PDF file path or file-like stream."""
    try:
        from pypdf import PdfReader
        
        reader = PdfReader(file_source)
        pages_text = []
        for i, page in enumerate(reader.pages):
            text = page.extract_text()
            if text:
                pages_text.append(text.strip())
        return "\n\n".join(pages_text)
    except Exception as e:
        raise ValueError(f"Failed to extract text from PDF: {str(e)}")

def extract_text_from_docx(file_source: Union[str, Path, BinaryIO]) -> str:
    """Extract clean text content from a DOCX file path or file-like stream."""
    try:
        import docx
        doc = docx.Document(file_source)
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        
        # Also extract table cells text
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    paragraphs.append(row_text)
                    
        return "\n".join(paragraphs)
    except Exception as e:
        raise ValueError(f"Failed to extract text from DOCX: {str(e)}")

def extract_text_from_txt(file_source: Union[str, Path, BinaryIO]) -> str:
    """Extract text from plain text file or bytes stream."""
    try:
        if isinstance(file_source, (str, Path)):
            with open(file_source, "r", encoding="utf-8", errors="replace") as f:
                return f.read()
        elif hasattr(file_source, "read"):
            content = file_source.read()
            if isinstance(content, bytes):
                return content.decode("utf-8", errors="replace")
            return content
        else:
            return str(file_source)
    except Exception as e:
        raise ValueError(f"Failed to extract text from text source: {str(e)}")

def extract_text(file_source: Union[str, Path, BinaryIO], filename: str = "") -> str:
    """
    Universal extractor: determines file format from path or provided filename
    and returns parsed plain text.
    """
    if isinstance(file_source, (str, Path)):
        path = Path(file_source)
        suffix = path.suffix.lower()
        if not path.exists():
            raise FileNotFoundError(f"File not found: {file_source}")
    else:
        suffix = Path(filename).suffix.lower() if filename else ".txt"

    if suffix == ".pdf":
        return extract_text_from_pdf(file_source)
    elif suffix in (".docx", ".doc"):
        return extract_text_from_docx(file_source)
    else:
        return extract_text_from_txt(file_source)
